#v2
import psutil
import win32gui
import win32process
import time
from openpyxl import load_workbook
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_PATH = os.path.join(SCRIPT_DIR, "..", "data", "hours.xlsx")
GENERIC_TITLES = {None, "Roblox Studio", "Auto-Recovery"}

parent_pid = None
last_title = None
log_name = None
active = False
target_row = None

start_time = 0

wb = None
sheet = None

def find_parent() -> (int | None):
    for p in psutil.process_iter(['pid', 'name']):
        if p.info['name'] == "RobloxStudioBeta.exe":
            return p.info['pid']
    return None

def get_window_titles_for_pid(target : str) -> list[str]:
    titles = []
    def callback(hwnd, _):
        if win32gui.IsWindowVisible(hwnd):
            _, pid = win32process.GetWindowThreadProcessId(hwnd)
            if pid == target:
                title = win32gui.GetWindowText(hwnd)
                if title: titles.append(title)
        return True

    win32gui.EnumWindows(callback, None)
    return titles

def prompt_log_name():
    global log_name, target_row
    log_name = input("Enter the project you would like to log hours for:\n")

    target_row = None
    for row_num in range(1, sheet.max_row + 1):
        if sheet.cell(row=row_num, column=1).value == log_name:
            target_row = row_num
            break

    if target_row is None:
        target_row = sheet.max_row + 1
        sheet.cell(row=target_row, column=1, value=log_name)
        sheet.cell(row=target_row, column=2, value=0.0)

def resolve_to_hours(time_ns : int) -> float:
    return time_ns / 1000000000 / 3600

while True:
    if parent_pid is None:
        parent_pid = find_parent()
        if parent_pid:
            print("Setting parent pid:", parent_pid)
            if wb is None:
                wb = load_workbook(DATA_PATH)
                sheet = wb.active
                if sheet is None:
                    wb.save(DATA_PATH)
                    wb.close()
                    wb = None; continue;
                if wb is None: continue;
        
        time.sleep(1)
        continue

    if not psutil.pid_exists(parent_pid):
        print("Studio closed")
        if active:
            et = resolve_to_hours(time.monotonic_ns() - start_time)
            current_hrs = sheet.cell(row=target_row, column=2).value
            sheet.cell(row=target_row, column=2, value=current_hrs + et)

        if wb is not None:
            wb.save(DATA_PATH)
            wb.close()

        parent_pid = None
        last_title = None
        active = False
        wb = None
        sheet = None

        time.sleep(1)
        continue

    titles  = get_window_titles_for_pid(parent_pid)
    current = titles[0] if titles else None

    if current != last_title:
        is_generic = current in GENERIC_TITLES
        was_generic = last_title in GENERIC_TITLES

        if was_generic and not is_generic:
            active = True
            print("Experience loaded, starting log:", current)
            prompt_log_name()

            start_time = time.monotonic_ns()

        elif not was_generic and is_generic:
            active = False

            et = resolve_to_hours(time.monotonic_ns() - start_time)
            current_hrs = sheet.cell(row=target_row, column=2).value
            sheet.cell(row=target_row, column=2, value=current_hrs+et)
            wb.save(DATA_PATH)

            log_name = None
            target_row = None

            print("Experience closed, stopping log")

        last_title = current

    

    time.sleep(1)